"""Convert the source lab PDF to native Word content and append an editable project."""
import logging
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import json
from copy import deepcopy
from pdf2docx import Converter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxcompose.composer import Composer
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
TEMP=ROOT/'tmp/docx-lab-record'
OUTPUT=ROOT/'output/docx/AML_Lab_Record_BatteryTwin_Editable.docx'
SOURCE=Path(r'C:\Users\SriRamSasi\OneDrive\Desktop\AML LAB.pdf')
COVER=Path(r'C:\Users\SriRamSasi\OneDrive\Pictures\Screenshots\Screenshot 2026-10-14 163515.png')

ABSTRACT=('BatteryTwin is an AI-based battery-health monitoring project developed using 636 measured discharge cycles from four NASA lithium-ion laboratory cells. '
          'The system validates capacity and sensor measurements, creates causal historical features, and compares Linear Regression, Random Forest, XGBoost, LSTM and GRU models for next-cycle State of Health forecasting. '
          'Battery-wise training, validation and testing reduce information leakage; Optuna tuning, PCA, SHAP and error analysis support reproducible evaluation. '
          'On 160 held-out forecasts, persistence achieves a lower error than the learned SOH models. Exploratory remaining-life models, novelty detection, derived charge tracking and utilization planning are presented through an editable-input Streamlit dashboard. '
          'The local application passes 37 automated tests. Docker and Hugging Face deployment packaging are prepared, while container execution and public hosting remain unverified. '
          'The project demonstrates an interpretable monitoring workflow; the small laboratory dataset does not establish real EV pack safety or physical fault diagnosis.')

def base_doc():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=Pt(48);sec.bottom_margin=Pt(45)
    sec.left_margin=Pt(55);sec.right_margin=Pt(55)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(6)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=Pt(13.5)
    doc.styles['Title'].font.size=Pt(14);doc.styles['Title'].font.bold=True
    doc.styles['Title'].paragraph_format.space_after=Pt(8)
    for name in ['Heading 1','Heading 2']:
        doc.styles[name].font.size=Pt(12);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(8)
        doc.styles[name].paragraph_format.space_after=Pt(5)
    return doc

def border(section):
    old=section._sectPr.find(qn('w:pgBorders'))
    if old is not None: section._sectPr.remove(old)
    borders=OxmlElement('w:pgBorders');borders.set(qn('w:offsetFrom'),'page')
    for side in ['top','left','bottom','right']:
        edge=OxmlElement(f'w:{side}')
        for key,value in [('val','single'),('sz','6'),('space','24'),('color','000000')]: edge.set(qn(f'w:{key}'),value)
        borders.append(edge)
    section._sectPr.append(borders)

def text(doc,content='',bold=False,size=None,align=None,after=None):
    para=doc.add_paragraph();run=para.add_run(content);run.bold=bold
    if size:run.font.size=Pt(size)
    if align is not None:para.alignment=align
    if after is not None:para.paragraph_format.space_after=Pt(after)
    return para

def heading(doc,content):doc.add_heading(content,level=1)

def make_cover():
    doc=base_doc();sec=doc.sections[0];border(sec)
    sec.top_margin=Pt(92)
    p=doc.add_paragraph('COIMBATORE INSTITUTE OF TECHNOLOGY','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    text(doc,'(Government-Aided Autonomous Institution Affiliated to Anna University, Chennai)',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=10)
    text(doc,'COIMBATORE 14',size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=38)
    text(doc,'ADVANCED MACHINE LEARNING LAB - 19MAM56',bold=True,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
    text(doc,'SUBMITTED TO',bold=True,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=8)
    text(doc,'Ms.Hemashree.P',bold=True,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=8)
    text(doc,'Dr.K.Swathine',bold=True,size=12,align=WD_ALIGN_PARAGRAPH.CENTER,after=22)
    with Image.open(COVER) as im:
        logo=im.crop((234,399,405,573));logo.save(TEMP/'cit-logo.png')
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(TEMP/'cit-logo.png'),width=Inches(2.05))
    p.paragraph_format.space_after=Pt(35)
    details=[('SUBMITTED BY',True),('Name: SriRam Sasi M',False),('Reg No: 2403717674421052',False),
             ('Department: M.Sc. AIML',False),('Date: 14 October 2026',False),
             ('Subject: Advanced Machine Learning Lab',False),('Subject Code: 19MAM56',False)]
    for label,bold in details:
        p=text(doc,label,bold=bold,size=11,after=5)
        p.paragraph_format.left_indent=Inches(3.1)
    doc.save(TEMP/'editable-cover.docx')
    return doc

def table(doc,rows,widths):
    tbl=doc.add_table(rows=0,cols=len(widths));tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
    for col,width in zip(tbl.columns,widths):col.width=Inches(width)
    props=tbl._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        tag=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','5'),('color','D9D9D9')]:tag.set(qn('w:'+key),value)
        borders.append(tag)
    props.append(borders)
    for idx,values in enumerate(rows):
        cells=tbl.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
            p.paragraph_format.line_spacing=Pt(11)
            if j>0 and len(values)>2:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            run=p.add_run(str(value));run.font.size=Pt(9.5);run.bold=idx==0
            if idx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(shade)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');tbl.rows[-1]._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');tbl.rows[-1]._tr.get_or_add_trPr().append(no_split)
    text(doc,'',after=1)

def picture(doc,name,caption,height):
    path=ROOT/'artifacts/plots'/name
    with Image.open(path) as im:w,h=im.size
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    maxw=6.5;maxh=height
    width=min(maxw,maxh*w/h)
    p.add_run().add_picture(str(path),width=Inches(width))
    p.paragraph_format.space_after=Pt(3)
    text(doc,caption,size=9,after=5)

def code(doc,lines):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=Pt(10)
    run=p.add_run(lines);run.font.name='Courier New';run.font.size=Pt(8.5)

def make_project():
    doc=base_doc();border(doc.sections[0])
    p=doc.add_paragraph('PROJECT REPORT','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('BatteryTwin AI Based EV Battery Health Monitoring','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading(doc,'ABSTRACT');text(doc,ABSTRACT)
    heading(doc,'AIM')
    text(doc,'To implement battery-health monitoring using measured discharge histories to forecast next-cycle SOH, estimate useful cycles remaining, and flag unusual patterns. Models are compared against simple baselines and presented in a Streamlit application.')
    heading(doc,'PROBLEM STATEMENT')
    text(doc,'Repeated use reduces battery capacity. BatteryTwin combines past capacity and sensor measurements to show degradation and support exploratory maintenance planning. EV monitoring motivates the project; the experiment uses laboratory single-cell data.')
    heading(doc,'DATASET AND TARGETS')
    text(doc,'NASA supplies 636 discharge cycles from B0005, B0006, B0007 and B0018. SOH = capacity / 2 Ah x 100. EOL is the first observed capacity at or below 1.4 Ah. Inputs include voltage, current, temperature, elapsed discharge time and capacity.')
    heading(doc,'ALGORITHM')
    steps=['Extract and validate NASA discharge measurements.',
           'Create lag, rolling-capacity and sensor features using completed history only.',
           'Train on B0005/B0007, validate on B0018 and test on B0006.',
           'Fit training-only imputation, scaling and PCA; create battery-specific sequences.',
           'Train models, tune with Optuna, evaluate common targets and explain errors with SHAP.',
           'Save artifacts and serve validated predictions through Streamlit.']
    for i,step in enumerate(steps,1):text(doc,f'{i}. {step}',after=3)
    doc.add_page_break()
    p=doc.add_paragraph('IMPLEMENTATION','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading(doc,'Data processing and feature engineering')
    text(doc,'The loader extracts actual MATLAB fields into cycle summaries. Validation checks required fields, sensor ranges and duplicate identifiers. IQR outliers are reported and retained. Lag features include previous capacity and previous voltage/current/temperature summaries; rolling statistics describe completed cycles. Imputation and scaling use training batteries only.')
    heading(doc,'Modeling tuning and explanation')
    text(doc,'Linear Regression, Random Forest, PCA + Linear, XGBoost, LSTM and GRU forecast SOH. Recurrent models use Adam, early stopping, learning-rate scheduling and gradient clipping. Optuna executes eight XGBoost trials and four LSTM trials. Forest, LSTM and Gaussian-process models estimate RUL. Isolation Forest flags novelty; SHAP explains the SOH forest.')
    heading(doc,'Main implementation modules')
    table(doc,[['Module','Purpose'],['src/data/loader.py','Read measured discharge records'],
       ['src/data/feature_engineering.py','Build causal features and sequences'],['scripts/train.py','Train tune evaluate and save models'],
       ['scripts/evaluate.py','External baseline and error analysis'],['api/schemas.py and api/services.py','Validate inputs and run saved models'],
       ['streamlit_app.py','Interactive monitoring and predictions'],['Dockerfile','Package Streamlit and fitted artifacts']],[3.05,3.45])
    heading(doc,'Saved model prediction flow')
    text(doc,'Streamlit shares Pydantic validation and BatteryService with the existing FastAPI backend. Cached resources load saved models; custom inputs are validated before inference. Training is an offline operation.')
    code(doc,'request = HistoryRequest.model_validate(payload)\nservice = BatteryService(ROOT / "models")\nsoh = service.predict_soh(request)\nrul = service.predict_rul(request)\nanomaly = service.detect_anomaly(request)')
    doc.add_page_break()
    p=doc.add_paragraph('OUTPUT AND APPLICATION FEATURES','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    text(doc,'Seven Streamlit views cover Dashboard, Predictions, Charge Tracking, Life Planning, Analytics, Model Performance and About & Requirements. Users select a battery and cycle, edit measurements or upload JSON history, and inspect predictions and plots.')
    heading(doc,'Measured health history')
    picture(doc,'soh.png','Figure 1  Measured SOH relative to a 2 Ah rated capacity',2.6)
    heading(doc,'Predicted versus measured health')
    picture(doc,'prediction_actual.png','Figure 2  Predictions on common held out B0006 targets',2.6)
    heading(doc,'Implemented features')
    text(doc,'Measured SOH is shown separately from forecasts. RUL compares forest, recurrent and Gaussian estimates. SHAP contributions explain model behavior; PCA, correlations, trial histories and loss curves support analysis.')
    text(doc,'Charge Tracking estimates a derived discharge-relative proxy because NASA has no measured SOC. Life Planning translates a fixed remaining-cycle estimate into days and does not establish lifetime improvement.')
    doc.add_page_break()
    p=doc.add_paragraph('EVALUATION EXECUTION AND RESULT','Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    heading(doc,'Evaluation')
    text(doc,'Primary SOH models share 160 B0006 forecasts (cycles 9-168). SOH MAE/RMSE use percentage points; RUL errors use cycles. Model selection uses validation performance. Values below are from executed experiment files.')
    import csv
    with (ROOT/'artifacts/results/benchmark.csv').open(newline='',encoding='utf8') as f:bench=list(csv.DictReader(f))
    rows=[['Model','Task','MAE','RMSE','R squared']]
    for r in bench:rows.append([r['model'],r['task'].upper(),f"{float(r['mae']):.3f}",f"{float(r['rmse']):.3f}",f"{float(r['r2']):.3f}"])
    table(doc,rows,[2.35,.65,1.05,1.05,1.4])
    text(doc,'The adapted two-lag linear baseline achieves MAE 0.684. Persistence beats the learned SOH models in the primary benchmark. Error analysis covers early-life, near-EOL, high-temperature and large-residual cases.')
    heading(doc,'Execution and verification')
    code(doc,'.\\start-streamlit.ps1\n# Open http://127.0.0.1:8501\n.\\deploy-docker.ps1\n# Requires a working Docker Linux engine')
    text(doc,'The complete suite has 37 passing tests, including nine Streamlit view/form tests. The local app is verified. Docker packages a non-root Streamlit runtime on port 7860. Local container execution remains blocked by Docker startup failure. Hugging Face returned HTTP 403 because the current credential lacks Space-creation permission; no live cloud deployment is claimed.')
    heading(doc,'Result')
    text(doc,'BatteryTwin was implemented and verified locally for next-cycle health forecasting, exploratory RUL, novelty detection and visual analysis. Four laboratory cells do not establish EV pack safety. Anomalies lack physical-fault labels; nominal 90% Gaussian ranges cover only 26.25% of held-out observations. Larger datasets and calibrated uncertainty remain necessary.')
    heading(doc,'References')
    text(doc,'NASA Li-ion Battery Aging Datasets: catalog.data.gov/dataset/li-ion-battery-aging-datasets. The supplied EV repository and full literature references are documented in docs/references.json.',size=9.5)
    doc.save(TEMP/'editable-project.docx')
    return doc

def build():
    TEMP.mkdir(parents=True,exist_ok=True);OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    converted=TEMP/'editable-original-lab.docx'
    if not converted.exists():
        converter=Converter(str(SOURCE))
        try:converter.convert(str(converted),multi_processing=False)
        finally:converter.close()
    cover=make_cover();project=make_project();lab=Document(converted)
    # An explicit section break separates the front page from the converted lab pages.
    cover.add_section(WD_SECTION_START.NEW_PAGE)
    # Cover sections must not force the lab text to use the cover's large top margin.
    cover.sections[-1].top_margin=Pt(48)
    compose=Composer(cover);compose.append(lab)
    # A new page is inserted before the four-page native Word project section.
    compose.doc.add_page_break()
    compose.append(project)
    compose.doc.core_properties.author='SriRam Sasi M'
    compose.doc.core_properties.title='Advanced Machine Learning Lab Record and BatteryTwin Project'
    compose.save(OUTPUT)
    doc=Document(OUTPUT)
    alltext='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for required in ['Date: 14 October 2026','Subject: Advanced Machine Learning Lab','Subject Code: 19MAM56','ABSTRACT','636 measured discharge cycles','RANDOM FOREST CLASSIFIER']:
        assert required in alltext,required
    log={'output':str(OUTPUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),
         'abstract_words':len(ABSTRACT.split()),'lab_source':'79-page source PDF converted to native Word content',
         'cover':'Editable text and supplied logo','date':'14 October 2026','subject':'Advanced Machine Learning Lab - 19MAM56'}
    (TEMP/'structural-verification.json').write_text(json.dumps(log,indent=2),encoding='utf8')
    print(json.dumps(log,indent=2),flush=True)

if __name__=='__main__':build()

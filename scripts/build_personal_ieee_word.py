"""Create an editable Word report from the same report flowables as the PDF."""
import contextlib
import io
from html.parser import HTMLParser
from pathlib import Path
from types import ModuleType

from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import Paragraph, Image, Table, Spacer
from reportlab.graphics.shapes import Drawing

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'output/docx/BatteryTwin_IEEE_Project_Report_SriRam_Sasi.docx'


class RichText(HTMLParser):
    def __init__(self, paragraph):
        super().__init__(convert_charrefs=True)
        self.p = paragraph
        self.bold = self.italic = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('b', 'strong'): self.bold += 1
        if tag in ('i', 'em'): self.italic += 1
        if tag == 'br': self.p.add_run().add_break()

    def handle_endtag(self, tag):
        if tag in ('b', 'strong'): self.bold = max(0, self.bold - 1)
        if tag in ('i', 'em'): self.italic = max(0, self.italic - 1)

    def handle_data(self, data):
        r = self.p.add_run(data)
        r.bold = True if self.bold else None
        r.italic = True if self.italic else None


def source_story():
    source = ROOT / 'scripts/build_report.py'
    module = ModuleType('report_for_word')
    module.__file__ = str(source)
    # The renderer only needs the workspace root, not ML environment imports.
    code = source.read_text(encoding='utf-8').replace(
        'from src.utils.config import ROOT',
        'ROOT = Path(__file__).resolve().parents[1]')
    exec(compile(code, str(source), 'exec'), module.__dict__)
    captured = []
    class CaptureDocument:
        def __init__(self, *args, **kwargs): pass
        def addPageTemplates(self, *args, **kwargs): pass
        def build(self, story): captured.extend(story)
    module.BaseDocTemplate = CaptureDocument
    with contextlib.redirect_stdout(io.StringIO()):
        module.report(ROOT / 'tmp/pdfs/personal-ieee/ieee-paper.pdf')
    return captured


def section_setup(section, columns=1):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = section.right_margin = Pt(40)
    section.top_margin = section.bottom_margin = Pt(42)
    section.header_distance = section.footer_distance = Pt(20)
    cols = section._sectPr.find(qn('w:cols'))
    cols.set(qn('w:num'), str(columns))
    cols.set(qn('w:space'), '360')


def cover(doc):
    s = doc.sections[0]
    section_setup(s)
    s.top_margin = Pt(69)
    border = OxmlElement('w:pgBorders')
    border.set(qn('w:offsetFrom'), 'page')
    for side in ['top', 'left', 'bottom', 'right']:
        edge = OxmlElement('w:' + side)
        for key, value in [('val', 'single'), ('sz', '5'), ('space', '24'), ('color', '000000')]:
            edge.set(qn('w:' + key), value)
        border.append(edge)
    s._sectPr.append(border)
    def centered(text, size, bold=False, after=12, style=None):
        p = doc.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = Pt(size + 3)
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
        return p
    centered('COIMBATORE INSTITUTE OF TECHNOLOGY', 15, True, 23)
    centered('(Government-Aided Autonomous Institution Affiliated to Anna University, Chennai)', 11.5, after=18)
    centered('COIMBATORE - 641 014', 12, after=25)
    centered('ADVANCED MACHINE LEARNING LAB - 19MAM56', 12, True, 12)
    centered('CAT II PROJECT REPORT', 12, True, 19)
    centered('BatteryTwin', 23, True, 8, 'Title')
    centered('AI-Based EV Battery Health Monitoring', 15, True, 22)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(26)
    p.add_run().add_picture(str(ROOT / 'tmp/pdfs/lab-report/cit-logo.png'), width=Pt(130))
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Pt(250))
    p.paragraph_format.space_after = Pt(9)
    p.add_run('SUBMITTED TO\tSUBMITTED BY').bold = True
    rows = [('Ms. Hemashree.P', 'Name: SriRam Sasi M'),
            ('Dr. K. Swathine', 'Reg No: 2403717674421052'),
            ('', 'Department: M.Sc. AIML'), ('', 'Date: 21 October 2026'),
            ('', 'Subject: Advanced Machine Learning Lab'), ('', 'Subject Code: 19MAM56')]
    for left, right in rows:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Pt(250))
        p.paragraph_format.line_spacing = Pt(18)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(left + '\t' + right); r.font.size = Pt(10.5)


def add_flowable(doc, item):
    name = item.__class__.__name__
    if isinstance(item, Paragraph):
        style = item.style.name
        target_style = {'Head': 'Heading 1', 'Caption': 'Caption', 'Small': 'Small Text'}.get(style, 'Normal')
        p = doc.add_paragraph(style=target_style)
        RichText(p).feed(item.text)
        if style == 'Caption':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = item.getPlainText().startswith('TABLE')
        elif style == 'Head':
            p.paragraph_format.keep_with_next = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if style == 'Small' else WD_ALIGN_PARAGRAPH.JUSTIFY
    elif isinstance(item, Image):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(item.filename), width=Pt(item.drawWidth), height=Pt(item.drawHeight))
    elif isinstance(item, Table):
        table = doc.add_table(rows=len(item._cellvalues), cols=len(item._cellvalues[0]))
        table.autofit = False
        for col, width in zip(table.columns, item._colWidths): col.width = Pt(width)
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            e = OxmlElement('w:' + edge)
            for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]: e.set(qn('w:' + key), value)
            borders.append(e)
        table._tbl.tblPr.append(borders)
        for i, values in enumerate(item._cellvalues):
            for j, value in enumerate(values):
                cell = table.cell(i, j); cell.width = Pt(item._colWidths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins = OxmlElement('w:tcMar')
                for side, size in [('top','60'),('bottom','60'),('left','70'),('right','70')]:
                    margin = OxmlElement('w:' + side)
                    margin.set(qn('w:w'), size); margin.set(qn('w:type'), 'dxa'); margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(9)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(str(value)); run.font.size = Pt(7.2); run.bold = i == 0
                if i == 0:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EDEDED'); cell._tc.get_or_add_tcPr().append(shade)
            props = table.rows[i]._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if i == 0: props.append(OxmlElement('w:tblHeader'))
    elif isinstance(item, Drawing):
        for text in ['NASA measured discharge cycles', 'Validate schema and battery IDs',
                     'Causal lag and rolling features', 'Battery held out training and Optuna',
                     'Saved ML LSTM and GRU models', 'SHAP and Streamlit FastAPI interface']:
            p = doc.add_paragraph(text, style='Small Text')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
    elif isinstance(item, Spacer):
        # Paragraph spacing supplies the source's small inter-block gaps.
        pass
    elif name == '_FrameBreak':
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)
        p.add_run().add_break(WD_BREAK.COLUMN)
    elif name == 'PageBreak':
        doc.add_page_break()
    else:
        raise ValueError(f'Unsupported report element: {name}')


def build():
    story = source_story()
    doc = Document()
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'; style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(9)
    normal.paragraph_format.line_spacing = Pt(11.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = False
    for name, size, leading, before, after in [('Heading 1',10.2,13,8,7), ('Caption',7.5,9,0,8)]:
        s = doc.styles[name]; s.font.size = Pt(size)
        s.paragraph_format.line_spacing = Pt(leading)
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
    doc.styles['Heading 1'].font.bold = True
    small = doc.styles.add_style('Small Text', 1)
    small.base_style = normal; small.font.size = Pt(7.8)
    small.paragraph_format.line_spacing = Pt(10)
    doc.styles['Caption'].font.italic = True
    doc.styles['Caption'].font.bold = False
    compat = doc.settings.element.find(qn('w:compat'))
    if compat is None:
        compat = OxmlElement('w:compat'); doc.settings.element.append(compat)
    compat.append(OxmlElement('w:doNotExpandShiftReturn'))
    cover(doc)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section_setup(section)
    border = section._sectPr.find(qn('w:pgBorders'))
    if border is not None: section._sectPr.remove(border)
    title = doc.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run(story[0].getPlainText()); r.font.size = Pt(15); r.bold = True
    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.space_after = Pt(20)
    author.paragraph_format.line_spacing = Pt(12)
    RichText(author).feed(story[1].text)
    for r in author.runs: r.font.size = Pt(10)
    section = doc.add_section(WD_SECTION_START.CONTINUOUS)
    section_setup(section, 2)
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for item in story[3:]: add_flowable(doc, item)
    doc.core_properties.author = 'SriRam Sasi M'
    doc.core_properties.title = 'BatteryTwin IEEE Project Report'
    doc.core_properties.subject = '19MAM56 Advanced Machine Learning Lab CAT II'
    # Remove built-in theme font and title-rule residue from Word templates.
    for element in doc.styles.element.iter(qn('w:pBdr')):
        element.getparent().remove(element)
    for p in doc.paragraphs:
        for element in list(p._p.iter(qn('w:pBdr'))): element.getparent().remove(element)
        for run in p.runs: run.font.name = 'Times New Roman'
    for fonts in doc.styles.element.iter(qn('w:rFonts')):
        for name in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:
            fonts.attrib.pop(qn('w:' + name), None)
        for name in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:' + name), 'Times New Roman')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    return OUTPUT


if __name__ == '__main__': build()
